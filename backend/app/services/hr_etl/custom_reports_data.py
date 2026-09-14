"""Custom reports — list and export views from the tenant custom_reports schema."""
from __future__ import annotations

import io
import json
import re
import time
from dataclasses import dataclass, field
from datetime import datetime, timezone
from decimal import Decimal
from typing import Any, Literal

from fpdf import FPDF
from fpdf.fonts import FontFace
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import text
from sqlalchemy.engine import Engine

from app.services.hr_etl.payroll_report_period_config import (
    payroll_period_filter_spec,
    payroll_period_columns_to_omit,
    supports_payroll_period_filter,
)
from app.services.hr_etl.schema_names import custom_reports_schema

_VIEW_NAME_RE = re.compile(r"^[a-z][a-z0-9_]*$", re.IGNORECASE)
_COLUMN_NAME_RE = re.compile(r"^[a-z][a-z0-9_]*$", re.IGNORECASE)
ExportFormat = Literal["xlsx", "pdf"]
_EXPORT_BRAND_FOOTER = "Powered by MintHRM"
_EXCEL_HEADER_FONT = Font(bold=True, color="1E293B")
_EXCEL_HEADER_FILL = PatternFill(fill_type="solid", fgColor="F1F5F9")
_EXPORT_OMIT_COLUMNS = frozenset({"tenant_id"})
_FILTER_OMIT_COLUMNS = frozenset({"tenant_id"})
_MAX_FILTER_COLUMNS = 12
_MAX_VALUES_PER_COLUMN = 50
_MAX_DISTINCT_OPTIONS = 200
_VIEW_COLUMNS_CACHE_TTL_SEC = 300
_view_columns_cache: dict[tuple[str, str], tuple[float, list[str]]] = {}
_PAYSLIP_LINE_COLUMNS = ("item_name", "unit", "amount")
_PAYSLIP_ORDER_BY = "proc_year DESC, proc_month DESC, emp_no, sort_order, section, item_name"

@dataclass
class ExportFilters:
    proc_year: int | None = None
    proc_month: int | None = None
    emp_no: str | None = None
    column_filters: dict[str, list[str]] = field(default_factory=dict)

    def has_period(self) -> bool:
        return self.proc_year is not None and self.proc_month is not None

    def has_column_filters(self) -> bool:
        return any(values for values in self.column_filters.values())


_REPORT_LABEL_ACRONYMS = {
    "ot": "OT",
    "je": "JE",
}
_EXPORT_COLUMN_LABEL_OVERRIDES = {
    "proc_year": "Year",
    "proc_month": "Month",
    "reporting_year": "Year",
    "reporting_month": "Month",
}
_NON_AMOUNT_COLUMNS = frozenset({
    "proc_year",
    "proc_month",
    "proc_half",
    "reporting_year",
    "reporting_month",
    "year",
    "month",
    "month_number",
    "emp_no",
    "employee_number",
    "sort_order",
    "headcount",
    "count_of_emp_no",
    "nopay_days",
    "line_count",
    "event_count",
    "active_headcount",
    "net_active_headcount_change_mom",
})
_AMOUNT_COLUMN_RE = re.compile(
    r"(?:^|_)(amount|salary|cost|rate|pay|tax|deduction|allowance|premium|"
    r"balance|remittance|contribution|incentive|pegged|hours|gross|net|"
    r"deductions|payment|payable|ot(?:_|$)|sum_of_)",
    re.IGNORECASE,
)
_EXCEL_AMOUNT_NUMBER_FORMAT = "#,##0.00"


def _format_label_word(part: str) -> str:
    lower = part.lower()
    if lower in _REPORT_LABEL_ACRONYMS:
        return _REPORT_LABEL_ACRONYMS[lower]
    return lower[:1].upper() + lower[1:] if lower else ""


def _label_from_view_name(view_name: str) -> str:
    name = view_name
    if name.startswith("vw_"):
        name = name[3:]
    return " ".join(
        _format_label_word(part)
        for part in name.replace("_", " ").split()
        if part
    )


def _custom_report_display_label(
    view_name: str,
    stored_name: str | None = None,
) -> str:
    """Prefer catalog report_name; fall back to derived label from view_name."""
    if stored_name and stored_name.strip():
        return stored_name.strip()
    return _label_from_view_name(view_name)


def _resolve_export_label(
    tenant_id: str,
    view_name: str,
    *,
    pg: Engine | None = None,
) -> str:
    """Report title for exports: configured report_name when available."""
    from app.services.hr_etl.custom_reports_catalog import get_report_definition_sync

    definition = get_report_definition_sync(tenant_id, view_name, pg=pg)
    stored = definition.report_name if definition else None
    return _custom_report_display_label(view_name, stored)


def _list_views_from_information_schema(
    pg: Engine, tenant_id: str
) -> list[dict[str, str | bool]]:
    schema = custom_reports_schema(tenant_id)
    with pg.connect() as conn:
        rows = conn.execute(
            text(
                """
                SELECT table_name
                FROM information_schema.views
                WHERE table_schema = :schema
                ORDER BY table_name
                """
            ),
            {"schema": schema},
        ).scalars().all()

    return [_custom_report_view_item(name, schema, tenant_id) for name in rows]


def _custom_report_view_item(
    view_name: str,
    schema: str,
    tenant_id: str,
    *,
    label: str | None = None,
    module: str = "payroll",
    source: str = "sync",
    is_payslip: bool | None = None,
) -> dict[str, str | bool]:
    payslip = (
        is_payslip_view(view_name, tenant_id)
        if is_payslip is None
        else is_payslip
    )
    return {
        "view_name": view_name,
        "label": _custom_report_display_label(view_name, label),
        "schema": schema,
        "module": module,
        "source": source,
        "is_payslip": payslip,
        "supports_period_filter": supports_payroll_period_filter(
            view_name, tenant_id=tenant_id
        ),
    }


def list_custom_report_views(
    pg: Engine,
    tenant_id: str,
    *,
    module: str | None = None,
) -> list[dict[str, str | bool]]:
    from app.services.hr_etl.custom_reports_catalog import load_report_definitions_sync

    schema = custom_reports_schema(tenant_id)
    definitions = load_report_definitions_sync(tenant_id, module=module)
    if not definitions:
        views = _list_views_from_information_schema(pg, tenant_id)
        if module and module.strip():
            key = module.strip().lower()
            views = [view for view in views if str(view.get("module", "")).lower() == key]
        return views

    return [
        _custom_report_view_item(
            row.view_name,
            schema,
            tenant_id,
            label=row.report_name,
            module=row.module,
            source=row.source,
            is_payslip=row.is_payslip,
        )
        for row in definitions
    ]


def _view_exists(pg: Engine, schema: str, view_name: str) -> bool:
    with pg.connect() as conn:
        found = conn.execute(
            text(
                """
                SELECT 1
                FROM information_schema.views
                WHERE table_schema = :schema
                  AND table_name = :view_name
                LIMIT 1
                """
            ),
            {"schema": schema, "view_name": view_name},
        ).scalar()
    return bool(found)


def parse_column_filters_json(raw: str | None) -> dict[str, list[str]]:
    if not raw or not raw.strip():
        return {}
    try:
        data = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise ValueError("Invalid column_filters JSON") from exc
    if not isinstance(data, dict):
        raise ValueError("column_filters must be a JSON object")
    parsed: dict[str, list[str]] = {}
    for key, values in data.items():
        if not isinstance(key, str) or not _COLUMN_NAME_RE.match(key):
            raise ValueError(f"Invalid filter column name: {key!r}")
        if not isinstance(values, list) or not values:
            continue
        cleaned = [str(v) for v in values if v is not None and str(v).strip() != ""]
        if cleaned:
            parsed[key.lower()] = cleaned[:_MAX_VALUES_PER_COLUMN]
    if len(parsed) > _MAX_FILTER_COLUMNS:
        raise ValueError(f"At most {_MAX_FILTER_COLUMNS} column filters allowed")
    return parsed


def invalidate_view_columns_cache(tenant_id: str | None = None) -> None:
    if tenant_id is None:
        _view_columns_cache.clear()
        return
    prefix = f"{tenant_id.strip()}:"
    for key in list(_view_columns_cache):
        if key[0].startswith(prefix):
            del _view_columns_cache[key]


def _get_view_columns(pg: Engine, schema: str, view_name: str, *, tenant_id: str) -> list[str]:
    cache_key = (f"{tenant_id}:{schema}", view_name)
    now = time.monotonic()
    cached = _view_columns_cache.get(cache_key)
    if cached and now - cached[0] < _VIEW_COLUMNS_CACHE_TTL_SEC:
        return list(cached[1])

    with pg.connect() as conn:
        rows = conn.execute(
            text(
                """
                SELECT column_name
                FROM information_schema.columns
                WHERE table_schema = :schema
                  AND table_name = :view_name
                ORDER BY ordinal_position
                """
            ),
            {"schema": schema, "view_name": view_name},
        ).scalars().all()
    columns = list(rows)
    _view_columns_cache[cache_key] = (now, columns)
    return columns


def _filterable_columns(
    columns: list[str],
    *,
    view_name: str | None = None,
    tenant_id: str | None = None,
) -> list[str]:
    omit = {name.lower() for name in _FILTER_OMIT_COLUMNS}
    if view_name:
        omit.update(
            payroll_period_columns_to_omit(view_name, tenant_id=tenant_id)
        )
    return [
        col
        for col in columns
        if col.lower() not in omit and _COLUMN_NAME_RE.match(col)
    ]


def _validate_column_filters(
    column_filters: dict[str, list[str]],
    valid_columns: set[str],
) -> dict[str, list[str]]:
    valid_lower = {col.lower(): col for col in valid_columns}
    validated: dict[str, list[str]] = {}
    for col, values in column_filters.items():
        actual = valid_lower.get(col.lower())
        if not actual:
            raise ValueError(f"Column {col!r} is not filterable for this report")
        validated[actual] = values
    return validated


def _build_view_where(
    view_name: str,
    filters: ExportFilters,
    valid_columns: set[str],
    *,
    tenant_id: str | None = None,
) -> tuple[str, dict[str, Any]]:
    where_parts: list[str] = []
    params: dict[str, Any] = {}

    if filters.has_period():
        period_spec = payroll_period_filter_spec(view_name, tenant_id=tenant_id)
        if is_payslip_view(view_name, tenant_id):
            year_column = "proc_year"
            month_column = "proc_month"
        elif period_spec is not None:
            year_column = period_spec.year_column
            month_column = period_spec.month_column
        else:
            raise ValueError("Period filter is not supported for this report")
        if year_column not in valid_columns or month_column not in valid_columns:
            raise ValueError(
                f"Period columns {year_column!r} and {month_column!r} "
                f"are not available for report {view_name!r}"
            )
        where_parts.append(f'"{year_column}" = :proc_year')
        where_parts.append(f'"{month_column}" = :proc_month')
        params["proc_year"] = filters.proc_year
        params["proc_month"] = filters.proc_month
    if filters.emp_no:
        if not is_payslip_view(view_name, tenant_id):
            raise ValueError("Employee filter is only supported for payslip export")
        where_parts.append("emp_no = :emp_no")
        params["emp_no"] = filters.emp_no

    if filters.column_filters:
        validated = _validate_column_filters(filters.column_filters, valid_columns)
        for col, values in validated.items():
            placeholders: list[str] = []
            for index, value in enumerate(values):
                key = f"cf_{col}_{index}"
                placeholders.append(f":{key}")
                params[key] = value
            quoted = f'"{col}"'
            where_parts.append(f"{quoted}::text IN ({', '.join(placeholders)})")

    where_sql = f" WHERE {' AND '.join(where_parts)}" if where_parts else ""
    return where_sql, params


def list_column_filter_options(
    pg: Engine, tenant_id: str, view_name: str, *, column: str | None = None
) -> dict[str, Any]:
    schema, _ = _resolve_view_context(pg, tenant_id, view_name)
    if is_payslip_view(view_name, tenant_id):
        raise ValueError("Column filters are not available for payslip reports")

    all_columns = _get_view_columns(pg, schema, view_name, tenant_id=tenant_id)
    columns = _filterable_columns(
        all_columns, view_name=view_name, tenant_id=tenant_id
    )[:_MAX_FILTER_COLUMNS]

    if column is None:
        return {
            "view_name": view_name,
            "columns": [{"column": col, "values": []} for col in columns],
        }

    if not _COLUMN_NAME_RE.match(column):
        raise ValueError(f"Invalid filter column name: {column!r}")
    valid_lower = {col.lower(): col for col in columns}
    actual = valid_lower.get(column.lower())
    if not actual:
        raise ValueError(f"Column {column!r} is not filterable for this report")

    with pg.connect() as conn:
        result = conn.execute(
            text(
                f"""
                SELECT DISTINCT "{actual}"::text AS value
                FROM "{schema}"."{view_name}"
                WHERE "{actual}" IS NOT NULL
                ORDER BY value
                LIMIT :limit
                """
            ),
            {"limit": _MAX_DISTINCT_OPTIONS},
        )
        values = [row[0] for row in result.fetchall()]

    return {"view_name": view_name, "columns": [{"column": actual, "values": values}]}


def _exclude_columns(
    columns: list[str],
    rows: list[tuple[Any, ...]],
    omit: frozenset[str],
) -> tuple[list[str], list[tuple[Any, ...]]]:
    if not omit:
        return columns, rows
    omit_lower = {name.lower() for name in omit}
    keep = [i for i, col in enumerate(columns) if col.lower() not in omit_lower]
    if len(keep) == len(columns):
        return columns, rows
    return [columns[i] for i in keep], [tuple(row[i] for i in keep) for row in rows]


def _fetch_view_page(
    pg: Engine,
    schema: str,
    view_name: str,
    *,
    tenant_id: str | None = None,
    limit: int,
    offset: int,
    filters: ExportFilters | None = None,
) -> tuple[list[str], list[tuple[Any, ...]], int]:
    filters = filters or ExportFilters()
    tid = tenant_id or ""
    valid_columns = set(_get_view_columns(pg, schema, view_name, tenant_id=tid))
    where_sql, params = _build_view_where(
        view_name, filters, valid_columns, tenant_id=tenant_id
    )
    order_by = ""

    with pg.connect() as conn:
        total = conn.execute(
            text(f'SELECT COUNT(*) FROM "{schema}"."{view_name}"{where_sql}'),
            params,
        ).scalar_one()
        result = conn.execute(
            text(
                f'SELECT * FROM "{schema}"."{view_name}"{where_sql}{order_by}'
                f" LIMIT :limit OFFSET :offset"
            ),
            {**params, "limit": limit, "offset": offset},
        )
        columns = list(result.keys())
        rows = [tuple(row) for row in result.fetchall()]
    return columns, rows, int(total)


def _json_cell(value: Any) -> Any:
    if value is None:
        return None
    if hasattr(value, "isoformat"):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, (bool, int, float, str)):
        return value
    return str(value)


def is_payslip_view(view_name: str, tenant_id: str | None = None) -> bool:
    if tenant_id:
        from app.services.hr_etl.custom_reports_catalog import get_report_definition_sync

        definition = get_report_definition_sync(tenant_id, view_name)
        if definition is not None:
            return definition.is_payslip
    return False


def _resolve_view_context(
    pg: Engine, tenant_id: str, view_name: str
) -> tuple[str, str]:
    _validate_view_name(view_name)
    schema = custom_reports_schema(tenant_id)
    if not _view_exists(pg, schema, view_name):
        raise LookupError(f"Report view not found: {view_name}")
    return schema, view_name


def _row_to_dict(columns: list[str], row: tuple[Any, ...]) -> dict[str, Any]:
    return {
        col: _json_cell(value) for col, value in zip(columns, row, strict=True)
    }


def _column_index(columns: list[str], name: str) -> int:
    lowered = {col.lower(): idx for idx, col in enumerate(columns)}
    try:
        return lowered[name.lower()]
    except KeyError as exc:
        raise ValueError(f"Expected column {name!r} in payslip view") from exc


def list_payslip_filters(pg: Engine, tenant_id: str, view_name: str) -> dict[str, Any]:
    schema, _ = _resolve_view_context(pg, tenant_id, view_name)
    if not is_payslip_view(view_name, tenant_id):
        raise ValueError(f"View {view_name!r} is not a payslip report")

    with pg.connect() as conn:
        result = conn.execute(
            text(
                f"""
                SELECT DISTINCT
                    proc_year,
                    proc_month,
                    emp_no,
                    emp_name,
                    emp_full_name,
                    designation_name
                FROM "{schema}"."{view_name}"
                ORDER BY proc_year DESC, proc_month DESC, emp_no
                """
            )
        )
        columns = list(result.keys())
        entries = [_row_to_dict(columns, row) for row in result.fetchall()]

    return {"view_name": view_name, "entries": entries}


def preview_payslip(
    pg: Engine,
    tenant_id: str,
    view_name: str,
    *,
    emp_no: str,
    proc_year: int,
    proc_month: int,
) -> dict[str, Any]:
    schema, _ = _resolve_view_context(pg, tenant_id, view_name)
    if not is_payslip_view(view_name, tenant_id):
        raise ValueError(f"View {view_name!r} is not a payslip report")

    with pg.connect() as conn:
        result = conn.execute(
            text(
                f"""
                SELECT *
                FROM "{schema}"."{view_name}"
                WHERE emp_no = :emp_no
                  AND proc_year = :proc_year
                  AND proc_month = :proc_month
                ORDER BY sort_order, section, item_name
                """
            ),
            {"emp_no": emp_no, "proc_year": proc_year, "proc_month": proc_month},
        )
        columns = list(result.keys())
        rows = [tuple(row) for row in result.fetchall()]

    columns, rows = _exclude_columns(columns, rows, _EXPORT_OMIT_COLUMNS)
    if not rows:
        raise LookupError(
            f"No payslip found for employee {emp_no!r} in {proc_year}-{proc_month:02d}"
        )

    first = _row_to_dict(columns, rows[0])
    line_columns = [col for col in _PAYSLIP_LINE_COLUMNS if col in columns]
    line_rows = [
        {col: row_dict.get(col) for col in line_columns}
        for row_dict in (_row_to_dict(columns, row) for row in rows)
    ]

    return {
        "mode": "payslip",
        "view_name": view_name,
        "label": _label_from_view_name(view_name),
        "schema": schema,
        "employee": {
            "emp_no": first.get("emp_no"),
            "emp_name": first.get("emp_name"),
            "emp_full_name": first.get("emp_full_name"),
            "designation_name": first.get("designation_name"),
        },
        "period": {"proc_year": proc_year, "proc_month": proc_month},
        "columns": line_columns,
        "rows": line_rows,
        "total": len(line_rows),
    }


def preview_view_data(
    pg: Engine,
    tenant_id: str,
    view_name: str,
    *,
    limit: int = 50,
    offset: int = 0,
    filters: ExportFilters | None = None,
) -> dict[str, Any]:
    _validate_view_name(view_name)
    schema = custom_reports_schema(tenant_id)
    if not _view_exists(pg, schema, view_name):
        raise LookupError(f"Report view not found: {view_name}")

    page_limit = min(max(limit, 1), 200)
    page_offset = max(offset, 0)
    filters = filters or ExportFilters()
    columns, rows, total = _fetch_view_page(
        pg,
        schema,
        view_name,
        tenant_id=tenant_id,
        limit=page_limit,
        offset=page_offset,
        filters=filters,
    )
    columns, rows = _exclude_columns(columns, rows, _EXPORT_OMIT_COLUMNS)
    row_dicts = [
        {col: _json_cell(value) for col, value in zip(columns, row, strict=True)}
        for row in rows
    ]
    return {
        "view_name": view_name,
        "label": _label_from_view_name(view_name),
        "schema": schema,
        "columns": columns,
        "rows": row_dicts,
        "total": total,
        "limit": page_limit,
        "offset": page_offset,
        "applied_filters": filters.column_filters,
    }


def _validate_view_name(view_name: str) -> None:
    if not _VIEW_NAME_RE.match(view_name):
        raise ValueError(f"Invalid view name: {view_name!r}")


def _load_view_export_data(
    pg: Engine,
    tenant_id: str,
    view_name: str,
    filters: ExportFilters | None = None,
) -> tuple[str, list[str], list[tuple[Any, ...]]]:
    schema, _ = _resolve_view_context(pg, tenant_id, view_name)
    filters = filters or ExportFilters()
    valid_columns = set(_get_view_columns(pg, schema, view_name, tenant_id=tenant_id))
    where_sql, params = _build_view_where(
        view_name, filters, valid_columns, tenant_id=tenant_id
    )
    order_by = (
        f" ORDER BY {_PAYSLIP_ORDER_BY}"
        if is_payslip_view(view_name, tenant_id)
        else ""
    )

    with pg.connect() as conn:
        result = conn.execute(
            text(f'SELECT * FROM "{schema}"."{view_name}"{where_sql}{order_by}'),
            params,
        )
        columns = list(result.keys())
        rows = [tuple(row) for row in result.fetchall()]

    if filters.has_period() and not rows:
        period_label = f"{filters.proc_year}-{filters.proc_month:02d}"
        raise LookupError(f"No data for period {period_label}")

    columns, rows = _exclude_columns(columns, rows, _EXPORT_OMIT_COLUMNS)
    return schema, columns, rows


def _export_filename(
    tenant_id: str,
    view_name: str,
    extension: str,
    filters: ExportFilters | None = None,
    *,
    pg: Engine | None = None,
) -> str:
    """Build a download filename from the report display label."""
    label = _resolve_export_label(tenant_id, view_name, pg=pg)
    safe = re.sub(r"[^\w\s-]", "", label)
    safe = re.sub(r"\s+", "_", safe.strip())
    if not safe:
        safe = view_name
    filters = filters or ExportFilters()
    if filters.has_period():
        safe = f"{safe}_{filters.proc_year}-{filters.proc_month:02d}"
    if filters.emp_no:
        emp = re.sub(r"[^\w-]", "", str(filters.emp_no))
        safe = f"{safe}_{emp}"
    if filters.has_column_filters():
        safe = f"{safe}_filtered"
    ext = extension.lstrip(".")
    return f"{safe}.{ext}"


def _is_amount_column(column: str) -> bool:
    lower = column.lower()
    if lower in _NON_AMOUNT_COLUMNS:
        return False
    if re.match(r"^(count|days_|event_|headcount|nopay_days)", lower):
        return False
    return bool(_AMOUNT_COLUMN_RE.search(lower))


def _coerce_number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        return float(value)
    return None


def _format_amount_display(value: Any) -> str:
    number = _coerce_number(value)
    if number is None:
        return _format_cell(value)
    return f"{number:,.2f}"


def _format_cell(value: Any) -> str:
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat(sep=" ", timespec="seconds")
    return str(value)


_PDF_CHAR_REPLACEMENTS = str.maketrans(
    {
        "\u2013": "-",  # en dash
        "\u2014": "-",  # em dash
        "\u2212": "-",  # minus sign
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u00a0": " ",
        "\u2026": "...",
    }
)


def _pdf_cell(value: Any, column: str | None = None) -> str:
    """Render text safely for core PDF fonts (latin-1 only)."""
    if column and _is_amount_column(column):
        text = _format_amount_display(value)
    else:
        text = _format_cell(value)
    text = text.translate(_PDF_CHAR_REPLACEMENTS)
    return text.encode("latin-1", errors="replace").decode("latin-1")


def _excel_cell(value: Any) -> Any:
    if value is None:
        return None
    if hasattr(value, "isoformat"):
        return value.isoformat(sep=" ", timespec="seconds")
    number = _coerce_number(value)
    if number is not None:
        return number
    return value


def _export_column_label(column: str) -> str:
    """Spreadsheet column heading: underscores to spaces, each word title-cased."""
    override = _EXPORT_COLUMN_LABEL_OVERRIDES.get(column.lower())
    if override:
        return override
    return " ".join(
        _format_label_word(part)
        for part in column.replace("_", " ").split()
        if part
    )


def _export_generated_at() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")


class _MintHrmExportPdf(FPDF):
    """PDF export with branded footer on every page."""

    def __init__(self, generated_at: str, *args: Any, **kwargs: Any) -> None:
        super().__init__(*args, **kwargs)
        self._generated_at = generated_at

    def footer(self) -> None:
        self.set_y(-10)
        self.set_font("Helvetica", "I", 7)
        self.set_text_color(100, 100, 100)
        self.cell(
            0,
            8,
            _pdf_cell(f"Generated: {self._generated_at}  |  {_EXPORT_BRAND_FOOTER}"),
            align="C",
        )


def export_view_to_excel(
    pg: Engine,
    tenant_id: str,
    view_name: str,
    filters: ExportFilters | None = None,
) -> tuple[bytes, str]:
    _schema, columns, rows = _load_view_export_data(
        pg, tenant_id, view_name, filters=filters
    )
    if is_payslip_view(view_name, tenant_id):
        columns, rows = _exclude_columns(columns, rows, frozenset({"section"}))

    label = _resolve_export_label(tenant_id, view_name, pg=pg)
    generated_at = _export_generated_at()

    wb = Workbook()
    ws = wb.active
    ws.title = label[:31]
    ws.append([label])
    ws.cell(row=1, column=1).font = Font(bold=True, size=12)
    ws.append([f"Generated: {generated_at}"])
    ws.append([])
    header_row = ws.max_row + 1
    for col_idx, column in enumerate(columns, start=1):
        cell = ws.cell(
            row=header_row,
            column=col_idx,
            value=_export_column_label(column),
        )
        cell.font = _EXCEL_HEADER_FONT
        cell.fill = _EXCEL_HEADER_FILL
    for row in rows:
        row_idx = ws.max_row + 1
        for col_idx, (column, value) in enumerate(zip(columns, row, strict=True), start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=_excel_cell(value))
            if _is_amount_column(column) and _coerce_number(value) is not None:
                cell.number_format = _EXCEL_AMOUNT_NUMBER_FORMAT
    ws.append([])
    footer_row = ws.max_row + 1
    ws.append([_EXPORT_BRAND_FOOTER])
    ws.cell(row=footer_row, column=1).font = Font(italic=True, color="64748B")

    buf = io.BytesIO()
    wb.save(buf)
    filename = _export_filename(tenant_id, view_name, "xlsx", filters, pg=pg)
    return buf.getvalue(), filename


def _pdf_column_label(column: str) -> str:
    return _pdf_cell(_export_column_label(column))


def _write_pdf_table(
    pdf: FPDF,
    columns: list[str],
    rows: list[tuple[Any, ...]],
) -> None:
    if not columns:
        return
    col_width = pdf.epw / len(columns)
    col_widths = tuple(col_width for _ in columns)
    head_style = FontFace(
        emphasis="BOLD",
        fill_color=(241, 245, 249),
        color=(30, 41, 59),
        size_pt=8,
    )
    pdf.set_font("Helvetica", "", 7)
    with pdf.table(
        width=pdf.epw,
        col_widths=col_widths,
        headings_style=head_style,
        first_row_as_headings=True,
        line_height=4.5,
        text_align="LEFT",
    ) as table:
        header = table.row()
        for column in columns:
            header.cell(_pdf_column_label(column), style=head_style)
        for row in rows:
            data_row = table.row()
            for column, value in zip(columns, row, strict=True):
                data_row.cell(_pdf_cell(value, column))


def _iter_payslip_employee_groups(
    columns: list[str], rows: list[tuple[Any, ...]]
) -> list[tuple[dict[str, Any], list[tuple[Any, ...]]]]:
    idx_year = _column_index(columns, "proc_year")
    idx_month = _column_index(columns, "proc_month")
    idx_emp_no = _column_index(columns, "emp_no")
    idx_emp_name = _column_index(columns, "emp_name")
    idx_emp_full = _column_index(columns, "emp_full_name")
    idx_desig = _column_index(columns, "designation_name")
    line_indices = [_column_index(columns, col) for col in _PAYSLIP_LINE_COLUMNS if col in columns]

    groups: list[tuple[dict[str, Any], list[tuple[Any, ...]]]] = []
    current_key: tuple[Any, ...] | None = None
    current_meta: dict[str, Any] | None = None
    current_lines: list[tuple[Any, ...]] = []

    for row in rows:
        key = (row[idx_year], row[idx_month], row[idx_emp_no])
        if key != current_key:
            if current_meta is not None:
                groups.append((current_meta, current_lines))
            current_key = key
            current_meta = {
                "proc_year": row[idx_year],
                "proc_month": row[idx_month],
                "emp_no": row[idx_emp_no],
                "emp_name": row[idx_emp_name],
                "emp_full_name": row[idx_emp_full],
                "designation_name": row[idx_desig],
            }
            current_lines = [tuple(row[i] for i in line_indices)]
        else:
            current_lines.append(tuple(row[i] for i in line_indices))

    if current_meta is not None:
        groups.append((current_meta, current_lines))
    return groups


def export_payslip_pdf(
    pg: Engine,
    tenant_id: str,
    view_name: str,
    filters: ExportFilters | None = None,
) -> tuple[bytes, str]:
    _schema, columns, rows = _load_view_export_data(
        pg, tenant_id, view_name, filters=filters
    )
    label = _resolve_export_label(tenant_id, view_name, pg=pg)
    generated_at = _export_generated_at()
    line_columns = [col for col in _PAYSLIP_LINE_COLUMNS if col in columns]
    groups = _iter_payslip_employee_groups(columns, rows)

    pdf = _MintHrmExportPdf(generated_at, orientation="P", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=16)
    pdf.set_margins(12, 12, 12)

    for index, (meta, line_rows) in enumerate(groups):
        pdf.add_page()
        if index == 0:
            pdf.set_font("Helvetica", "B", 14)
            pdf.cell(0, 8, _pdf_cell(label), new_x="LMARGIN", new_y="NEXT")
            pdf.set_font("Helvetica", "", 8)
            pdf.cell(0, 5, _pdf_cell(f"Generated: {generated_at}"), new_x="LMARGIN", new_y="NEXT")
            pdf.ln(2)

        period_label = f"{meta['proc_year']}-{int(meta['proc_month']):02d}"
        pdf.set_font("Helvetica", "B", 11)
        pdf.cell(
            0,
            7,
            _pdf_cell(
                f"{meta.get('emp_full_name') or meta.get('emp_name') or meta['emp_no']} "
                f"({meta['emp_no']})"
            ),
            new_x="LMARGIN",
            new_y="NEXT",
        )
        pdf.set_font("Helvetica", "", 9)
        pdf.cell(
            0,
            5,
            _pdf_cell(
                f"Period: {period_label}  |  Designation: {meta.get('designation_name') or '-'}"
            ),
            new_x="LMARGIN",
            new_y="NEXT",
        )
        pdf.ln(2)
        _write_pdf_table(pdf, line_columns, line_rows)

    if not groups:
        pdf.add_page()
        pdf.set_font("Helvetica", "I", 10)
        pdf.cell(0, 6, "No payslip data available.", new_x="LMARGIN", new_y="NEXT")

    buf = io.BytesIO()
    pdf.output(buf)
    filename = _export_filename(tenant_id, view_name, "pdf", filters, pg=pg)
    return buf.getvalue(), filename


def export_view_to_pdf(
    pg: Engine,
    tenant_id: str,
    view_name: str,
    filters: ExportFilters | None = None,
) -> tuple[bytes, str]:
    if is_payslip_view(view_name, tenant_id):
        return export_payslip_pdf(pg, tenant_id, view_name, filters=filters)

    _schema, columns, rows = _load_view_export_data(
        pg, tenant_id, view_name, filters=filters
    )
    label = _resolve_export_label(tenant_id, view_name, pg=pg)
    generated_at = _export_generated_at()

    pdf = _MintHrmExportPdf(generated_at, orientation="L", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=14)
    pdf.set_margins(8, 8, 8)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 8, _pdf_cell(label), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 8)
    pdf.cell(0, 5, _pdf_cell(f"Generated: {generated_at}"), new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 5, _pdf_cell(f"Rows: {len(rows):,}"), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)

    if not columns:
        pdf.set_font("Helvetica", "I", 10)
        pdf.cell(0, 6, "No columns in view.", new_x="LMARGIN", new_y="NEXT")
    else:
        _write_pdf_table(pdf, columns, rows)

    buf = io.BytesIO()
    pdf.output(buf)
    filename = _export_filename(tenant_id, view_name, "pdf", filters, pg=pg)
    return buf.getvalue(), filename


def export_view(
    pg: Engine,
    tenant_id: str,
    view_name: str,
    export_format: ExportFormat = "xlsx",
    filters: ExportFilters | None = None,
) -> tuple[bytes, str]:
    if export_format == "pdf":
        return export_view_to_pdf(pg, tenant_id, view_name, filters=filters)
    return export_view_to_excel(pg, tenant_id, view_name, filters=filters)
