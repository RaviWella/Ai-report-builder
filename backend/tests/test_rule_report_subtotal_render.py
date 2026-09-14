"""A group subtotal row (report_service._inject_subtotals tags rows
"_row_kind": "subtotal") must render distinctly (bold) in both Excel and PDF,
and must never break a report that doesn't use subtotals at all."""

from __future__ import annotations

import io

import openpyxl

from app.domain.report_spec import PresentationSpec
from app.query_engine.runner import QueryResult
from app.rendering.excel_renderer import render_excel
from app.rendering.pdf_renderer import _env, render_pdf


def _load_sheet(data: bytes):
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    return wb.active


def _meal_plan_result() -> QueryResult:
    return QueryResult(
        columns=["Employee", "Date", "Count"],
        rows=[
            {"Employee": "E1", "Date": "2026-06-01", "Count": 1},
            {"Employee": "E1", "Date": "Total", "Count": 1, "_row_kind": "subtotal"},
            {"Employee": "E2", "Date": "2026-06-01", "Count": 2},
        ],
        sql="", row_count=3, truncated=False,
    )


def test_excel_subtotal_row_is_bold_and_data_row_is_not():
    data = render_excel(_meal_plan_result(), PresentationSpec(title="Meal Plan"))
    ws = _load_sheet(data)

    subtotal_row = next(row for row in ws.iter_rows() if row[1].value == "Total")
    data_row = next(row for row in ws.iter_rows()
                     if row[1].value == "2026-06-01" and row[0].value == "E1")

    assert subtotal_row[0].font.bold is True
    assert data_row[0].font.bold is not True


def test_excel_export_without_subtotals_is_unaffected():
    # No "_row_kind" on any row (the normal case, e.g. Golden Key OT) — must
    # render exactly as before, nothing bolded.
    result = QueryResult(
        columns=["Employee", "Total Hours"],
        rows=[{"Employee": "E1", "Total Hours": 10}],
        sql="", row_count=1, truncated=False,
    )
    data = render_excel(result, PresentationSpec(title="OT"))
    ws = _load_sheet(data)
    data_row = next(row for row in ws.iter_rows() if row[0].value == "E1")
    assert data_row[0].font.bold is not True


def test_pdf_template_marks_subtotal_row_with_css_class():
    html = _env.get_template("report.html").render(
        title="Meal Plan", header=None, footer=None, logo_url=None, orientation="A4",
        columns=[{"label": "Employee"}, {"label": "Date"}],
        rows=[
            {"cells": [{"value": "E1", "css": ""}, {"value": "2026-06-01", "css": ""}], "subtotal": False},
            {"cells": [{"value": "E1", "css": ""}, {"value": "Total", "css": ""}], "subtotal": True},
        ],
        totals=None, provenance="",
    )
    assert '<tr class="subtotal">' in html
    assert '<tr class="">' in html


def test_pdf_export_renders_without_crashing_with_a_subtotal_row():
    pdf_bytes = render_pdf(_meal_plan_result(), PresentationSpec(title="Meal Plan"))
    assert pdf_bytes[:4] == b"%PDF"


def test_pdf_formats_a_bare_time_value_as_hh_mm():
    import datetime as _dt

    from app.rendering.pdf_renderer import _fmt

    assert _fmt(_dt.time(14, 5, 0), None) == "14:05"


def test_pdf_formats_a_timezone_aware_datetime_without_the_utc_offset():
    import datetime as _dt

    from app.rendering.pdf_renderer import _fmt

    tz = _dt.timezone(_dt.timedelta(hours=5, minutes=30))
    aware = _dt.datetime(2026, 6, 6, 6, 16, tzinfo=tz)
    assert _fmt(aware, None) == "2026-06-06 06:16"
