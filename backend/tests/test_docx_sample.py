"""Word .docx sample letters are read as flowing text (merge fields kept as «Name»)."""

from __future__ import annotations

import io

from app.ingestion.document_parser import _is_docx, extract_lines, is_text_document


def _docx_bytes(*paragraphs: str) -> bytes:
    from docx import Document

    doc = Document()
    for t in paragraphs:
        doc.add_paragraph(t)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_is_a_text_letter_sample():
    assert _is_docx("offer.docx", "")
    assert is_text_document("offer.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert not is_text_document("sheet.xlsx", "")


def test_extract_lines_reads_docx_paragraphs_and_merge_marks():
    content = _docx_bytes(
        "Confidential",
        "Dear «Display_Name»,",
        "Acceptance of Resignation",
    )
    lines = extract_lines(content, "letter.docx", "")
    assert "Confidential" in lines
    assert any("Display_Name" in ln for ln in lines)
    assert any("Acceptance of Resignation" in ln for ln in lines)


def _empty_mergefield_docx(*names: str) -> bytes:
    """A paragraph with one empty (no cached result) MERGEFIELD per name —
    field code only, matching how Word stores an unfilled merge placeholder."""
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document()
    p = doc.add_paragraph()
    for name in names:
        r = p.add_run()
        instr = r._r.makeelement(qn("w:instrText"), {})
        instr.text = f" MERGEFIELD {name} \\* MERGEFORMAT "
        r._r.append(instr)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_mergefield_dedup_does_not_use_substring_match():
    """Two empty MERGEFIELDs where one name is a substring of the other
    (Name ⊂ FullName) must BOTH survive — not have the second silently
    dropped by a naive substring-containment dedup check."""
    lines = extract_lines(_empty_mergefield_docx("FullName", "Name"), "letter.docx", "")
    joined = " ".join(lines)
    assert "«FullName»" in joined
    assert "«Name»" in joined
