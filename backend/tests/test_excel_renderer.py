"""Excel export formatting — dates must render as dates, not raw serials."""

from __future__ import annotations

import io
from datetime import date, datetime, timedelta, timezone

import openpyxl

from app.domain.report_spec import ColumnPresentation, PresentationSpec
from app.query_engine.runner import QueryResult
from app.rendering.excel_renderer import render_excel


def _load_sheet(data: bytes):
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    return wb.active


def test_excel_export_applies_date_format():
    result = QueryResult(
        columns=["Date of Birth"],
        rows=[
            {"Date of Birth": date(1995, 5, 10)},
            {"Date of Birth": date(1993, 5, 18)},
        ],
        sql="",
        row_count=2,
        truncated=False,
    )
    data = render_excel(result, PresentationSpec(title="Employees"))
    ws = _load_sheet(data)

    # Locate the data rows by VALUE (not a hardcoded cell ref) — the exact row
    # offset depends on branding-header rows that aren't this test's concern.
    # openpyxl round-trips an Excel date serial as a datetime (Excel itself has
    # no pure date-only storage type), so compare the date parts, not equality
    # against a plain `date`.
    def _is_date(value, expected: date) -> bool:
        return isinstance(value, (date, datetime)) and value.year == expected.year \
            and value.month == expected.month and value.day == expected.day

    cell = next(c for row in ws.iter_rows() for c in row if _is_date(c.value, date(1995, 5, 10)))
    assert "yyyy" in cell.number_format.lower()
    assert any(_is_date(c.value, date(1993, 5, 18)) for row in ws.iter_rows() for c in row)


def test_excel_export_applies_datetime_format():
    result = QueryResult(
        columns=["Logged At"],
        rows=[{"Logged At": datetime(2024, 3, 15, 9, 30, 0)}],
        sql="",
        row_count=1,
        truncated=False,
    )
    data = render_excel(result, PresentationSpec(title="Audit"))
    ws = _load_sheet(data)

    # Locate the data row by VALUE — see the comment on the date-format test above.
    cell = next(c for row in ws.iter_rows() for c in row if c.value == datetime(2024, 3, 15, 9, 30, 0))
    assert "yyyy" in cell.number_format.lower()


def test_excel_export_applies_time_format():
    # A bare `time` column (e.g. a punch-in time) must render as HH:MM, not the
    # raw Excel time serial (a fraction of a day, e.g. 0.586...). Locate the data
    # row by VALUE (not a hardcoded cell ref) — the exact row offset depends on
    # branding-header rows that aren't this test's concern.
    from datetime import time

    result = QueryResult(
        columns=["Time In"],
        rows=[{"Time In": time(14, 5, 0)}],
        sql="",
        row_count=1,
        truncated=False,
    )
    data = render_excel(result, PresentationSpec(title="Attendance"))
    ws = _load_sheet(data)

    cell = next(c for row in ws.iter_rows() for c in row if c.value == time(14, 5, 0))
    assert "hh:mm" in cell.number_format.lower()


def test_excel_export_matches_a_reordered_presentation_column_order():
    # Real production bug: presentation.columns can be in a DIFFERENT order
    # than a row's own key order (the SQL SELECT — and so each row dict's key
    # order — follows dataSpec.fields' order, which the export's presentation-
    # column order doesn't always mirror, e.g. after reordering columns in the
    # builder). The renderer used to match a row's Nth value to the Nth
    # presentation column by bare POSITION, silently shifting every value into
    # the wrong column whenever the two orders diverged.
    result = QueryResult(
        columns=["Category", "Department", "In Time"],
        rows=[{"Category": "Executive", "Department": "Finance", "In Time": "08:00"}],
        sql="", row_count=1, truncated=False,
    )
    presentation = PresentationSpec(
        title="Test",
        columns=[
            ColumnPresentation(ref="employee.department", label="Department"),
            ColumnPresentation(ref="employee.category", label="Category"),
            ColumnPresentation(ref="attendance_daily.in_time", label="In Time"),
        ],
    )
    data = render_excel(result, presentation)
    ws = _load_sheet(data)
    header_row = next(row for row in ws.iter_rows(values_only=True) if row[0] == "Department")
    data_row = next(
        row for row in ws.iter_rows(values_only=True)
        if row[0] in ("Finance", "Executive", "08:00")
    )
    by_header = dict(zip(header_row, data_row, strict=True))
    assert by_header["Department"] == "Finance"
    assert by_header["Category"] == "Executive"
    assert by_header["In Time"] == "08:00"


def test_excel_export_strips_timezone_from_datetime_without_crashing():
    # A timestamptz straight from Postgres (e.g. punch_out_datetime) arrives
    # timezone-aware — xlsxwriter raises TypeError on that unless stripped first.
    tz = timezone(timedelta(hours=5, minutes=30))
    aware = datetime(2026, 6, 6, 6, 16, tzinfo=tz)
    result = QueryResult(
        columns=["Time Out"], rows=[{"Time Out": aware}], sql="", row_count=1, truncated=False,
    )
    data = render_excel(result, PresentationSpec(title="Shift Allowance"))   # must not raise
    ws = _load_sheet(data)
    cell = next(c for row in ws.iter_rows() for c in row if isinstance(c.value, datetime))
    assert cell.value == datetime(2026, 6, 6, 6, 16)   # tzinfo stripped, value unchanged otherwise
