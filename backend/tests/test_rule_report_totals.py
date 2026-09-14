"""A rule report's Excel/PDF grand-total row used to sum a HARDCODED pair of
column names (`{"ot_hours", "total_qualifying"}`, report_service.py) for
EVERY rule report's export — a platform-wide "no hardcoding" violation
(right for the one report those names came from, wrong or coincidentally-
wrong for any other tenant's rule report). `totals` is now an explicit,
per-report parameter threaded through create/update_rule_report into
presentation_spec, mirroring the existing row_number_column/subtotal
pattern (report_service.py's own `_rule_presentation` docstring already
documents these as "display concerns, never part of the governed
RuleReportSpec itself").
"""

from __future__ import annotations

from app.domain.rule_report import RuleReportSpec
from app.services.report_service import ReportService

_MIN_SPEC = {
    "name": "test_report",
    "source": "mart.mart_attendance_daily",
    "grain": ["employee_no"],
    "row_value": {"cases": [{"when": "worked_hours > 0", "value": "worked_hours"}], "else": "0"},
    "rollup": {"total": "sum(row_value)"},
    "output": ["employee_no", "total"],
}


def test_totals_defaults_to_empty_not_a_hardcoded_pair():
    rule = RuleReportSpec.model_validate(_MIN_SPEC)
    ps = ReportService._rule_presentation(rule, "Test Report", {})
    assert ps["totals"] == []


def test_totals_is_stored_verbatim_when_given():
    rule = RuleReportSpec.model_validate(_MIN_SPEC)
    ps = ReportService._rule_presentation(rule, "Test Report", {}, totals=["total"])
    assert ps["totals"] == ["total"]


def test_excel_totals_match_by_output_key_even_when_heading_is_renamed():
    """Rule reports flag totals by output field key; a renamed heading must not
    break the Excel totals row."""
    import io

    import openpyxl

    from app.domain.report_spec import ColumnPresentation, PresentationSpec
    from app.query_engine.runner import QueryResult
    from app.rendering.excel_renderer import render_excel

    result = QueryResult(
        columns=["employee_no", "total"],
        rows=[{"employee_no": "E1", "total": 10}, {"employee_no": "E2", "total": 20}],
        sql="",
        row_count=2,
        truncated=False,
    )
    presentation = PresentationSpec(
        title="Hours",
        columns=[
            ColumnPresentation(ref="employee_no", label="Employee No"),
            ColumnPresentation(ref="total", label="Total Individual Working Hours"),
        ],
    )
    data = render_excel(result, presentation, total_labels={"total"})
    ws = openpyxl.load_workbook(io.BytesIO(data), data_only=False).active
    total_row = next(
        row for row in ws.iter_rows(values_only=True) if row and row[0] == "Total"
    )
    assert total_row[1] is not None
    assert str(total_row[1]).startswith("=SUM(")
